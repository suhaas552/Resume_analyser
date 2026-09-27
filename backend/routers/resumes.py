import os
import uuid
from pathlib import Path

from docx import Document
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from pypdf import PdfReader

from database import get_db_connection
from routers.auth import get_current_user


router = APIRouter(
    prefix="/resumes",
    tags=["Resumes"]
)


# ---------------------------------------------------------
# Configuration
# ---------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent.parent
UPLOAD_DIR = BASE_DIR / "uploads"

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

ALLOWED_EXTENSIONS = {".pdf", ".docx"}

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB


# ---------------------------------------------------------
# Helper: Extract text from PDF
# ---------------------------------------------------------

def extract_text_from_pdf(file_path: Path) -> str:
    try:
        reader = PdfReader(str(file_path))

        extracted_pages = []

        for page in reader.pages:
            text = page.extract_text()

            if text:
                extracted_pages.append(text)

        return "\n".join(extracted_pages).strip()

    except Exception as error:
        raise RuntimeError(
            f"Could not extract text from PDF: {error}"
        )


# ---------------------------------------------------------
# Helper: Extract text from DOCX
# ---------------------------------------------------------

def extract_text_from_docx(file_path: Path) -> str:
    try:
        document = Document(str(file_path))

        paragraphs = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        # Also extract text from tables.
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()

                    if text:
                        paragraphs.append(text)

        return "\n".join(paragraphs).strip()

    except Exception as error:
        raise RuntimeError(
            f"Could not extract text from DOCX: {error}"
        )


# ---------------------------------------------------------
# Helper: Extract text according to extension
# ---------------------------------------------------------

def extract_resume_text(
    file_path: Path,
    extension: str
) -> str:

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)

    if extension == ".docx":
        return extract_text_from_docx(file_path)

    raise RuntimeError("Unsupported resume format")


# ---------------------------------------------------------
# Upload Resume
# ---------------------------------------------------------

@router.post(
    "/upload",
    status_code=status.HTTP_201_CREATED
)
async def upload_resume(
    file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user)
):
    if not file.filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No file was selected"
        )

    original_filename = Path(file.filename).name

    extension = Path(original_filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Only PDF and DOCX resume files are allowed"
        )

    # -----------------------------------------------------
    # Read uploaded file
    # -----------------------------------------------------

    file_content = await file.read()

    file_size = len(file_content)

    if file_size == 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Uploaded file is empty"
        )

    if file_size > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail="Resume file must not exceed 5 MB"
        )

    # -----------------------------------------------------
    # Generate safe unique filename
    # -----------------------------------------------------

    unique_name = f"{uuid.uuid4().hex}{extension}"

    stored_path = UPLOAD_DIR / unique_name

    relative_path = f"uploads/{unique_name}"

    # -----------------------------------------------------
    # Save file
    # -----------------------------------------------------

    try:
        with open(stored_path, "wb") as destination:
            destination.write(file_content)

    except Exception as error:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Could not save resume file: {error}"
        )

    connection = get_db_connection()

    if connection is None:
        # Remove file if database connection fails.
        if stored_path.exists():
            stored_path.unlink()

        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Database connection failed"
        )

    cursor = None
    resume_id = None

    try:
        cursor = connection.cursor(dictionary=True)

        # -------------------------------------------------
        # Create initial resume database record
        # -------------------------------------------------

        cursor.execute(
            """
            INSERT INTO resumes (
                user_id,
                original_filename,
                stored_filename,
                file_path,
                file_type,
                file_size,
                upload_status
            )
            VALUES (%s, %s, %s, %s, %s, %s, 'PROCESSING')
            """,
            (
                current_user["user_id"],
                original_filename,
                unique_name,
                relative_path,
                extension.replace(".", "").upper(),
                file_size
            )
        )

        resume_id = cursor.lastrowid

        connection.commit()

        # -------------------------------------------------
        # Extract resume text
        # -------------------------------------------------

        try:
            extracted_text = extract_resume_text(
                stored_path,
                extension
            )

            if not extracted_text:
                raise RuntimeError(
                    "No readable text was found in the resume"
                )

            # ---------------------------------------------
            # Mark processing as completed
            # ---------------------------------------------

            cursor.execute(
                """
                UPDATE resumes
                SET
                    extracted_text = %s,
                    upload_status = 'COMPLETED'
                WHERE resume_id = %s
                """,
                (
                    extracted_text,
                    resume_id
                )
            )

            connection.commit()

        except Exception as extraction_error:

            cursor.execute(
                """
                UPDATE resumes
                SET upload_status = 'FAILED'
                WHERE resume_id = %s
                """,
                (resume_id,)
            )

            connection.commit()

            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Resume uploaded but text extraction failed: {extraction_error}"
            )

        # -------------------------------------------------
        # Return saved resume
        # -------------------------------------------------

        cursor.execute(
            """
            SELECT
                resume_id,
                user_id,
                original_filename,
                stored_filename,
                file_path,
                file_type,
                file_size,
                upload_status,
                created_at,
                updated_at
            FROM resumes
            WHERE resume_id = %s
            """,
            (resume_id,)
        )

        saved_resume = cursor.fetchone()

        return {
            "message": "Resume uploaded and processed successfully",
            "resume": saved_resume
        }

    except HTTPException:
        raise

    except Exception as error:
        connection.rollback()

        if resume_id is None and stored_path.exists():
            stored_path.unlink()

        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Resume upload failed: {error}"
        )

    finally:
        if cursor is not None:
            cursor.close()

        if connection.is_connected():
            connection.close()


# ---------------------------------------------------------
# Get My Resumes
# ---------------------------------------------------------

@router.get("/my")
def get_my_resumes(
    current_user: dict = Depends(get_current_user)
):
    connection = get_db_connection()

    if connection is None:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Database connection failed"
        )

    cursor = None

    try:
        cursor = connection.cursor(dictionary=True)

        cursor.execute(
            """
            SELECT
                resume_id,
                original_filename,
                stored_filename,
                file_type,
                file_size,
                upload_status,
                created_at,
                updated_at
            FROM resumes
            WHERE user_id = %s
            ORDER BY created_at DESC
            """,
            (current_user["user_id"],)
        )

        resumes = cursor.fetchall()

        return {
            "count": len(resumes),
            "resumes": resumes
        }

    except Exception as error:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Could not retrieve resumes: {error}"
        )

    finally:
        if cursor is not None:
            cursor.close()

        if connection.is_connected():
            connection.close()


# ---------------------------------------------------------
# Get Single Resume
# ---------------------------------------------------------

@router.get("/{resume_id}")
def get_resume(
    resume_id: int,
    current_user: dict = Depends(get_current_user)
):
    connection = get_db_connection()

    if connection is None:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Database connection failed"
        )

    cursor = None

    try:
        cursor = connection.cursor(dictionary=True)

        cursor.execute(
            """
            SELECT
                resume_id,
                user_id,
                original_filename,
                stored_filename,
                file_type,
                file_size,
                extracted_text,
                upload_status,
                created_at,
                updated_at
            FROM resumes
            WHERE resume_id = %s
              AND user_id = %s
            """,
            (
                resume_id,
                current_user["user_id"]
            )
        )

        resume = cursor.fetchone()

        if resume is None:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Resume not found"
            )

        return {
            "resume": resume
        }

    except HTTPException:
        raise

    except Exception as error:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Could not retrieve resume: {error}"
        )

    finally:
        if cursor is not None:
            cursor.close()

        if connection.is_connected():
            connection.close()